# -*- coding: utf-8 -*-
"""Slips, printable schedules and the S-140."""
import io
from pathlib import Path

import pytest

import docx
import pypdf


def _text(pdf_bytes):
    return "\n".join(p.extract_text() for p in pypdf.PdfReader(io.BytesIO(pdf_bytes)).pages)


def _full_week(core, people, aux=True):
    weeks = {"W": {"parts": core.default_midweek_parts()}}
    slots = core.apply_aux(core.build_midweek_slots(weeks["W"]["parts"]), aux)
    names = dict(zip(core.get_students()["id"], core.get_students()["name"]))
    picks = {}
    for i, s in enumerate(slots):
        if s["role"] == "Aux Classroom Counselor":
            picks[i] = (people["Kofi Mensah"], None)
        elif s["student_part"]:
            aux_slot = s["hall"] == core.AUX_HALL
            picks[i] = (people["Efua Osei" if aux_slot else "Ama Owusu"],
                        people["Ama Owusu" if aux_slot else "Efua Osei"]
                        if s["needs_assistant"] else None)
        else:
            picks[i] = (people["Kofi Mensah"], None)
    core.save_schedule("2026-09-16", core.MIDWEEK, slots, picks,
                       {"heading": "SEPTEMBER 14–20", "aux": aux,
                        "opening_song": "Song 1"}, names)


def test_schedule_pdf_and_talk_title(core, people):
    _full_week(core, people)
    slots = core.default_weekend_slots()
    names = dict(zip(core.get_students()["id"], core.get_students()["name"]))
    core.save_schedule("2026-09-20", core.WEEKEND, slots,
                       {2: (people["Nii Tetteh"], None)},
                       {"talk_number": "12", "talk_title": "Is God Interested in You?"},
                       names)
    pdf = core.generate_schedule_pdf(
        [("2026-09-16", core.MIDWEEK), ("2026-09-20", core.WEEKEND)], core.get_schedules())
    text = _text(pdf)
    assert "Auxiliary classroom" in text
    assert "No. 12" in text and "Is God Interested in You?" in text


def test_schedule_pdf_in_ga(core, people):
    """Section headings switch language; the scripture reading, when saved,
    replaces the week label on the blue heading line."""
    _full_week(core, people)
    core.save_schedule("2026-09-16", core.MIDWEEK,
                       core.apply_aux(core.build_midweek_slots(
                           core.default_midweek_parts()), True),
                       {}, {"heading": "SEPTEMBER 14–20", "book": "YEREMIA 32-33",
                            "aux": True}, {})
    text_en = _text(core.generate_schedule_pdf(
        [("2026-09-16", core.MIDWEEK)], core.get_schedules()))
    assert "Treasures From God" in text_en
    assert "YEREMIA 32-33" in text_en          # the book line, not the week label
    assert "SEPTEMBER 14" not in text_en

    text_ga = _text(core.generate_schedule_pdf(
        [("2026-09-16", core.MIDWEEK)], core.get_schedules(), core.TRANSLATIONS["Ga"]))
    assert "Nyɔŋmɔ Wiemɔ" in text_ga
    assert "Treasures From God" not in text_ga
    assert "YEREMIA 32-33" in text_ga           # the reading itself is language-neutral


def test_s140_fills_both_halls(core, people, s140_template):
    _full_week(core, people)
    data, skipped = core.build_s140_data([("2026-09-16", core.MIDWEEK)],
                                         core.get_schedules(), "TEST CONG", "GROUP")
    assert not skipped and data["aux"] and data["clear_asa2"] is False
    week = data["weeks"][0]
    assert week["treasures"][2]["name2"] == "Efua Osei"
    assert week["aux_counselor"] == "Kofi Mensah"
    doc = docx.Document(io.BytesIO(core.fill_s140(s140_template, data)))
    cells = [c.text for r in doc.tables[0].rows for c in r.cells]
    assert "Efua Osei/Ama Owusu" in cells and "Ama Owusu/Efua Osei" in cells
    assert not any("0:00" == c for c in cells)


def test_a_second_auxiliary_classroom_flows_through(core):
    """The app only offers one auxiliary classroom, but nothing downstream
    should silently drop a second one if a schedule ever carries it."""
    slots = core.build_midweek_slots(core.default_midweek_parts())
    reading = next(s for s in slots if s["role"] == "Bible Reading")
    slots = slots + [
        {**reading, "hall": "aux_1"},
        {**reading, "hall": "aux_2"},
    ]
    picks = {len(slots) - 2: (None, None), len(slots) - 1: (None, None)}
    core.add_student("Adjeley Ayi", "Sister", ["Bible Reading"])
    core.add_student("Naa Koshie", "Sister", ["Bible Reading"])
    ids = dict(zip(core.get_students()["name"], core.get_students()["id"]))
    names = {v: k for k, v in ids.items()}
    picks = {len(slots) - 2: (ids["Adjeley Ayi"], None),
             len(slots) - 1: (ids["Naa Koshie"], None)}
    core.save_schedule("2026-09-16", core.MIDWEEK, slots, picks, {"aux": True}, names)

    rows = core.get_schedules()
    assert set(rows["hall"]) == {"main_hall", "aux_1", "aux_2"}

    # the slip data carries the third room through
    slip_rows = core.slip_rows_for(rows)
    assert {r["hall"] for r in slip_rows} >= {"aux_1", "aux_2"}

    # the part label names the room it is in
    label = core.slot_label({**reading, "hall": "aux_2"})
    assert "Auxiliary classroom 2" in label

    # and the S-140 pairing treats it as a classroom, not a duplicate main-hall part
    data, skipped = core.build_s140_data([("2026-09-16", core.MIDWEEK)], rows,
                                         "TEST CONG", "GROUP")
    assert not skipped and data["aux"]
    titles = [i["title"] for i in data["weeks"][0]["treasures"]]
    assert titles.count("Bible Reading") == 1      # not repeated once per classroom


def test_printed_schedule_is_fully_ga(core, people):
    """Interface stays English; everything on the printed sheet is Ga."""
    _full_week(core, people)
    rows = core.get_schedules()
    text = _text(core.generate_schedule_pdf(
        [("2026-09-16", core.MIDWEEK)], rows, core.TRANSLATIONS["Ga"]))

    # section headings, meeting name and room label all in Ga
    assert "Nyɔŋmɔ Wiemɔ" in text and "Hii Shi Akɛ Kristofonyo" in text
    assert "Wɔshiɛmɔ" in text
    assert "Asa 2" in text                         # the auxiliary classroom

    for english in ("Treasures From God", "Living as Christians",
                    "Midweek Meeting", "Auxiliary classroom"):
        assert english not in text, f"{english!r} still on the Ga sheet"


def test_weekend_sheet_is_ga(core, people):
    slots = core.default_weekend_slots()
    names = dict(zip(core.get_students()["id"], core.get_students()["name"]))
    core.save_schedule("2026-09-20", core.WEEKEND, slots,
                       {2: (people["Nii Tetteh"], None)}, {}, names)
    text = _text(core.generate_schedule_pdf(
        [("2026-09-20", core.WEEKEND)], core.get_schedules(), core.TRANSLATIONS["Ga"]))
    assert "Otsi Naagbee Kpee" in text
    assert "Weekend Meeting" not in text


def test_interface_labels_stay_english(core):
    """slot_label without hall names is the interface version."""
    slot = core.make_slot("Bible Reading", "Bible Reading", "Treasures",
                          3, 4, core.AUX_HALL)
    assert "Auxiliary classroom 1" in core.slot_label(slot)
    ga_rooms = {h: core.TRANSLATIONS["Ga"][h] for h in core.HALLS}
    assert "Asa 2" in core.slot_label(slot, ga_rooms)


def test_visitor_may_say_the_closing_prayer(core, people):
    """A guest speaker is often asked to close. The flag must also survive
    reopening a saved schedule, or the option vanishes on the next edit."""
    slots = core.default_weekend_slots()
    by_role = {(s["role"], s["title"]): s for s in slots}
    assert by_role[("Public Talk", "Public Talk Speaker")]["allow_visitor"]
    assert by_role[("Prayer", "Closing Prayer")]["allow_visitor"]
    assert not by_role[("Prayer", "Opening Prayer")]["allow_visitor"]
    assert not by_role[("Weekend Chairman", "Chairman")]["allow_visitor"]

    names = dict(zip(core.get_students()["id"], core.get_students()["name"]))
    closing = next(i for i, s in enumerate(slots)
                   if s["role"] == "Prayer" and s["title"] == "Closing Prayer")
    talk = next(i for i, s in enumerate(slots) if s["role"] == "Public Talk")
    core.save_schedule("2026-09-27", core.WEEKEND, slots,
                       {talk + 10000: "Bro. Addo — Osu",
                        closing + 10000: "Bro. Tetteh — Osu"}, {}, names)

    reopened, _, visitors = core.load_schedule("2026-09-27", core.WEEKEND,
                                               core.get_schedules())
    flags = {(s["role"], s["title"]): s["allow_visitor"] for s in reopened}
    assert flags[("Prayer", "Closing Prayer")]        # still offered on re-edit
    assert flags[("Public Talk", "Public Talk Speaker")]
    assert not flags[("Prayer", "Opening Prayer")]
    assert "Bro. Tetteh — Osu" in visitors.values()

    text = _text(core.generate_schedule_pdf([("2026-09-27", core.WEEKEND)],
                                            core.get_schedules()))
    assert "Bro. Tetteh — Osu" in text


def _weekend_with_guest(core, people):
    slots = core.default_weekend_slots()
    names = dict(zip(core.get_students()["id"], core.get_students()["name"]))
    talk = next(i for i, s in enumerate(slots) if s["role"] == "Public Talk")
    closing = next(i for i, s in enumerate(slots)
                   if s["role"] == "Prayer" and s["title"] == "Closing Prayer")
    picks = {i: (people["Kofi Mensah"], None) for i, s in enumerate(slots)}
    picks[talk] = (None, None)
    picks[talk + 10000] = "Bro. Addo — Osu"
    picks[closing] = (None, None)
    picks[closing + 10000] = "Bro. Tetteh — Osu"
    core.save_schedule("2026-09-27", core.WEEKEND, slots, picks,
                       {"talk_number": "12", "talk_title": "Is God Interested in You?"},
                       names)


def test_weekend_sheet_order_and_guest(core, people):
    """The closing prayer ends the sheet; a visitor is marked as a guest."""
    _weekend_with_guest(core, people)
    text = _text(core.generate_schedule_pdf([("2026-09-27", core.WEEKEND)],
                                            core.get_schedules()))
    for word in ("Chairman", "Public Talk", "Watchtower Study", "Theme",
                 "Guest speaker", "Opening Prayer", "Closing Prayer"):
        assert word in text, word
    assert text.index("Opening Prayer") < text.index("Public Talk")
    assert text.index("Public Talk") < text.index("Watchtower Study")
    assert text.index("Watchtower Study") < text.index("Closing Prayer")
    assert text.rstrip().index("Closing Prayer") > text.index("Chairman")
    assert "Bro. Tetteh — Osu" in text and "Bro. Addo — Osu" in text


def test_weekend_sheet_in_ga_has_no_english(core, people):
    _weekend_with_guest(core, people)
    text = _text(core.generate_schedule_pdf([("2026-09-27", core.WEEKEND)],
                                            core.get_schedules(),
                                            core.TRANSLATIONS["Ga"]))
    for word in ("Otsi Naagbee Kpee", "Sɛinɔtalɔ", "Maŋshiɛmɔ",
                 "Buu Mɔɔ Nikasemɔ", "Kanelɔ", "Saneyitso",
                 "Wielɔ ni afɔ lɛ nine", "Sɔlemɔ"):
        assert word in text, word
    for english in ("Chairman", "Public Talk", "Watchtower Study", "Theme",
                    "Guest speaker", "Closing Prayer", "Weekend Meeting"):
        assert english not in text, f"{english!r} leaked onto the Ga sheet"


def test_midweek_and_weekend_get_different_layouts(core, people):
    """One call, two sheets: the midweek running order and the weekend
    programme are laid out differently."""
    _full_week(core, people)
    _weekend_with_guest(core, people)
    text = _text(core.generate_schedule_pdf(
        [("2026-09-16", core.MIDWEEK), ("2026-09-27", core.WEEKEND)],
        core.get_schedules()))
    assert "Midweek Meeting" in text and "Weekend Meeting" in text
    assert "Treasures From God" in text          # midweek sections
    assert "Watchtower Study" in text            # weekend parts
    assert text.index("Midweek Meeting") < text.index("Weekend Meeting")


def test_auxiliary_classroom_group_is_printed(core, people):
    """Which group is using the classroom is stored per week and printed
    beside the counselor."""
    slots = core.apply_aux(core.build_midweek_slots(core.default_midweek_parts()), True)
    names = dict(zip(core.get_students()["id"], core.get_students()["name"]))
    counselor = next(i for i, s in enumerate(slots)
                     if s["role"] == "Aux Classroom Counselor")
    core.save_schedule("2026-09-16", core.MIDWEEK, slots,
                       {counselor: (people["Kofi Mensah"], None)},
                       {"aux": True, "aux_group": "1"}, names)
    assert core.get_meeting_meta("2026-09-16", core.MIDWEEK)["aux_group"] == "1"

    text = _text(core.generate_schedule_pdf([("2026-09-16", core.MIDWEEK)],
                                            core.get_schedules()))
    assert "Auxiliary classroom 1" in text and "Group 1" in text

    ga = _text(core.generate_schedule_pdf([("2026-09-16", core.MIDWEEK)],
                                          core.get_schedules(),
                                          core.TRANSLATIONS["Ga"]))
    assert "Asa 2" in ga and "Kuu 1" in ga
    assert "Group 1" not in ga

    # no group entered: the room is still named, with nothing after it
    core.save_schedule("2026-09-16", core.MIDWEEK, slots,
                       {counselor: (people["Kofi Mensah"], None)},
                       {"aux": True, "aux_group": ""}, names)
    plain = _text(core.generate_schedule_pdf([("2026-09-16", core.MIDWEEK)],
                                             core.get_schedules()))
    assert "Auxiliary classroom 1" in plain and "Group" not in plain


S89_BLANK = Path(__file__).resolve().parent / "tests_data" / "S-89_s-Mlt_GA.pdf"


@pytest.mark.skipif(not S89_BLANK.is_file(), reason="blank S-89 not in the repo")
def test_official_s89_is_filled_not_redrawn(core):
    """Printing on the real form gets its exact wording, so the app never has
    to carry a hand-typed copy of it."""
    blank = S89_BLANK.read_bytes()
    rows = [
        {"person": "Ɛfua Ɔsei", "assistant": "Naa Ŋmɛnɛ", "part_no": 4,
         "part_name": "x", "meeting_date": "2026-09-23", "hall": "aux_1"},
        {"person": "Samuel Oquaye", "assistant": None, "part_no": 5,
         "part_name": "x", "meeting_date": "2026-09-23", "hall": "aux_2"},
        {"person": "Vera Akitah", "assistant": None, "part_no": 3,
         "part_name": "x", "meeting_date": "2026-09-23", "hall": core.MAIN_HALL},
    ]
    text = _text(core.fill_s89(blank, rows))

    # the form's own wording, which the app does not store anywhere. Its Ga
    # labels cannot be read back by pypdf (the form's fonts carry no Unicode
    # mapping, as the workbook's do not), so check what does extract.
    assert "KPEE ASAIM" in text                  # the form's title
    assert "Asa 1" in text and "Asa 3" in text   # its room list
    # our values, Ga characters intact
    assert "Ɛfua Ɔsei" in text and "Naa Ŋmɛnɛ" in text
    assert "23 September 2026" in text
    assert text.count("X") == 3                  # one room ticked per filled slip

    # four slips to a page: three rows plus a blank still fills one sheet
    import pypdf, io
    assert len(pypdf.PdfReader(io.BytesIO(core.fill_s89(blank, rows))).pages) == 1
    assert len(pypdf.PdfReader(io.BytesIO(core.fill_s89(blank, rows * 2))).pages) == 2


@pytest.mark.skipif(not S89_BLANK.is_file(), reason="blank S-89 not in the repo")
@pytest.mark.skipif(not S89_BLANK.is_file(), reason="blank S-89 not in the repo")
def test_slips_need_the_official_blank(core):
    """The app no longer draws its own imitation of the form: without the blank
    it says so, rather than printing something that only looks official."""
    rows = [{"person": "Kofi Mensah", "assistant": None, "part_no": 3,
             "part_name": "x", "meeting_date": "2026-09-23", "hall": core.MAIN_HALL}]
    with pytest.raises(core.S89Error) as raised:
        core.slips_pdf(rows, core.TRANSLATIONS["Ga"], "Ga")
    assert "Admin" in str(raised.value)          # says where to put it

    core.save_template("s89_Ga", "S-89.pdf", S89_BLANK.read_bytes())
    official = _text(core.slips_pdf(rows, core.TRANSLATIONS["Ga"], "Ga"))
    assert "Kofi Mensah" in official and "KPEE ASAIM" in official

def test_midweek_column_widths(core):
    """One column set for every week. The part title needs the room the
    classroom's old second name column used to take."""
    page = 523.0
    widths = core.midweek_widths(page)
    assert len(widths) == 3
    assert abs(sum(widths) - page) < 1.0            # the full text width, no more
    assert widths[0] > page * 0.5                   # the title has room
    assert widths[-1] > widths[1]                   # names wider than their labels


def test_songs_print_in_the_sheet_language(core, people):
    """A Ga workbook stores "Lala 74" and an English one "Song 74"; the sheet
    prints whichever word matches it."""
    slots = core.build_midweek_slots(core.default_midweek_parts())
    names = dict(zip(core.get_students()["id"], core.get_students()["name"]))
    core.save_schedule("2026-09-16", core.MIDWEEK, slots, {},
                       {"opening_song": "Song 74", "closing_song": "Lala 134"}, names)
    rows = core.get_schedules()
    english = _text(core.generate_schedule_pdf([("2026-09-16", core.MIDWEEK)], rows))
    assert "Song 74" in english and "Song 134" in english
    assert "Lala" not in english
    ga = _text(core.generate_schedule_pdf([("2026-09-16", core.MIDWEEK)], rows,
                                          core.TRANSLATIONS["Ga"]))
    assert "Lala 74" in ga and "Lala 134" in ga
    assert "Song" not in ga


def test_guest_prayer_is_weekend_only(core):
    """A visiting speaker may close the weekend meeting; the midweek closing
    prayer is always a local brother."""
    weekend = {(s["role"], s["title"]): s for s in core.default_weekend_slots()}
    assert weekend[("Prayer", "Closing Prayer")]["allow_visitor"]

    midweek = {(s["role"], s["title"]): s for s in
               core.build_midweek_slots(core.default_midweek_parts())}
    assert not midweek[("Prayer", "Closing Prayer")]["allow_visitor"]
    assert not midweek[("Prayer", "Opening Prayer")]["allow_visitor"]


def test_s89_edge_cases(core):
    """Zero rows, one row, and a wrong file — the three ways the uploader and
    the print button can be used that are not the happy path."""
    import io

    import pypdf

    blank = S89_BLANK.read_bytes()

    # no assignments at all still prints one sheet of blanks, not zero pages
    empty = core.fill_s89(blank, [])
    assert len(pypdf.PdfReader(io.BytesIO(empty)).pages) == 1

    one = core.fill_s89(blank, [{"person": "Kofi Mensah", "assistant": None,
                                 "part_no": 3, "part_name": "x",
                                 "meeting_date": "2026-09-23",
                                 "hall": core.MAIN_HALL}])
    assert len(pypdf.PdfReader(io.BytesIO(one)).pages) == 1
    assert "Kofi Mensah" in _text(one)

    # checking a template does not render it
    assert core.check_s89_template(blank) == 4
    with pytest.raises(core.S89Error):
        core.check_s89_template(b"not a pdf at all")
    with pytest.raises(core.S89Error):
        core.check_s89_template(b"%PDF-1.4\n%%EOF\n")


def _midweek_week(core, date_iso, aux, people):
    slots = core.apply_aux(core.build_midweek_slots(core.default_midweek_parts()), aux)
    names = dict(zip(core.get_students()["id"], core.get_students()["name"]))
    picks = {i: (people["Kofi Mensah"],
                 people["Ama Owusu"] if s["needs_assistant"] else None)
             for i, s in enumerate(slots)}
    core.save_schedule(date_iso, core.MIDWEEK, slots, picks,
                       {"aux": aux, "aux_group": "1" if aux else "",
                        "opening_song": "Song 74", "middle_song": "Song 142",
                        "closing_song": "Song 134"}, names)


def _pages(pdf_bytes):
    import io

    import pypdf

    return len(pypdf.PdfReader(io.BytesIO(pdf_bytes)).pages)


def test_two_ordinary_weeks_share_a_sheet_at_full_size(core, people):
    """Compact pairs ordinary weeks without shrinking anything; a classroom
    week is too tall to pair, so it keeps its own sheet rather than being
    reduced to fit."""
    _midweek_week(core, "2026-09-09", False, people)
    _midweek_week(core, "2026-09-16", False, people)
    rows = core.get_schedules()
    both = [("2026-09-09", core.MIDWEEK), ("2026-09-16", core.MIDWEEK)]
    assert _pages(core.generate_schedule_pdf(both, rows, compact=True)) == 1
    assert _pages(core.generate_schedule_pdf(both, rows)) == 2      # off by default

    # the type is the same size either way
    one = core.generate_schedule_pdf([both[0]], rows)
    assert "Treasures From God" in _text(one)
    paired = _text(core.generate_schedule_pdf(both, rows, compact=True))
    assert paired.count("Treasures From God") == 2


def test_classroom_weeks_keep_their_own_sheet(core, people):
    _midweek_week(core, "2026-09-09", True, people)
    _midweek_week(core, "2026-09-16", True, people)
    rows = core.get_schedules()
    both = [("2026-09-09", core.MIDWEEK), ("2026-09-16", core.MIDWEEK)]
    assert _pages(core.generate_schedule_pdf(both, rows, compact=True)) == 2

    # mixed: an ordinary week cannot pair across a classroom week, because the
    # weeks have to stay in date order
    _midweek_week(core, "2026-09-23", False, people)
    _midweek_week(core, "2026-09-30", False, people)
    rows = core.get_schedules()
    mixed = [("2026-09-09", core.MIDWEEK), ("2026-09-16", core.MIDWEEK),
             ("2026-09-23", core.MIDWEEK), ("2026-09-30", core.MIDWEEK)]
    assert _pages(core.generate_schedule_pdf(mixed, rows, compact=True)) == 3


def test_classroom_has_its_own_section(core, people):
    """The classroom's parts are listed under a heading naming the room and
    group, not paired anonymously beside the main hall's names."""
    _midweek_week(core, "2026-09-16", True, people)
    rows = core.get_schedules()
    text = _text(core.generate_schedule_pdf([("2026-09-16", core.MIDWEEK)], rows))
    assert "Auxiliary classroom 1" in text and "Group 1" in text
    # the classroom's parts belong with the field ministry, so its heading sits
    # between that section and Living as Christians
    assert (text.index("Apply Yourself") < text.index("Auxiliary classroom 1")
            < text.index("Living as Christians"))

    ga = _text(core.generate_schedule_pdf([("2026-09-16", core.MIDWEEK)], rows,
                                          core.TRANSLATIONS["Ga"]))
    assert "Asa 2 – Kuu 1" in ga
    assert "Ŋaawolɔ" in ga                       # the counselor, listed once
    assert ga.count("Ŋaawolɔ") == 1


def test_watchtower_conductor_is_labelled_in_ga(core, people):
    slots = core.default_weekend_slots()
    names = dict(zip(core.get_students()["id"], core.get_students()["name"]))
    core.save_schedule("2026-09-20", core.WEEKEND, slots,
                       {i: (people["Kofi Mensah"], None) for i in range(len(slots))},
                       {}, names)
    ga = _text(core.generate_schedule_pdf([("2026-09-20", core.WEEKEND)],
                                          core.get_schedules(),
                                          core.TRANSLATIONS["Ga"]))
    assert "Buu Mɔɔ Nɔkwɛlɔ" in ga               # conductor, now labelled
    assert "Buu Mɔɔ Nikasemɔ" in ga              # the study itself


S140_EN = Path(__file__).resolve().parent / "tests_data" / "S-140_E.docx"
S140_GA = Path(__file__).resolve().parent / "tests_data" / "S-140_GA.docx"


@pytest.mark.skipif(not S140_EN.is_file(), reason="S-140 blanks not in the repo")
def test_both_published_s140_blanks_fill(core, people):
    """The published blank comes in each language and holds a single week. It
    used to be rejected unless it carried the Ga [DEETI] marker, and a month
    could not be exported from a one-week template."""
    import docx

    names = dict(zip(core.get_students()["id"], core.get_students()["name"]))
    dates = ["2026-09-09", "2026-09-16", "2026-09-23"]
    for d in dates:
        slots = core.build_midweek_slots(core.default_midweek_parts())
        picks = {i: (people["Kofi Mensah"],
                     people["Ama Owusu"] if s["needs_assistant"] else None)
                 for i, s in enumerate(slots)}
        core.save_schedule(d, core.MIDWEEK, slots, picks,
                           {"heading": f"WEEK OF {d}", "opening_song": "Song 74",
                            "middle_song": "Song 142", "closing_song": "Song 134"},
                           names)
    data, skipped = core.build_s140_data([(d, core.MIDWEEK) for d in dates],
                                         core.get_schedules(), "Teshie Asafo",
                                         "GROUP")
    assert not skipped and len(data["weeks"]) == 3

    for path, ga in ((S140_EN, False), (S140_GA, True)):
        blank = path.read_bytes()
        assert core.check_s140_template(blank) == 1      # one week in the blank
        out = core.fill_s140(blank, data)
        doc = docx.Document(io.BytesIO(out))
        text = "\n".join(c.text for r in doc.tables[0].rows for c in r.cells)
        # a block per week, each with its own heading
        for d in dates:
            assert f"WEEK OF {d}" in text
        assert "[DATE]" not in text and "[DEETI]" not in text
        assert "[Name]" not in text and "[Gbɛ́i]" not in text
        # the song word matches the template, not what was stored
        assert ("Lala 74" in text) is ga
        assert ("Song 74" in text) is (not ga)


@pytest.mark.skipif(not S140_EN.is_file(), reason="S-140 blanks not in the repo")
def test_s140_template_check_rejects_the_wrong_file(core):
    with pytest.raises(core.S140Error):
        core.check_s140_template(b"not a docx")
    with pytest.raises(core.S140Error):
        core.check_s140_template(S89_BLANK.read_bytes())      # a PDF


@pytest.mark.skipif(not S140_GA.is_file(), reason="S-140 blanks not in the repo")
def test_s140_fills_the_classroom_column(core, people):
    """The blank has an Asa 2 / Main Hall pair of name columns. The classroom's
    people belong in the first, and the counselor beside the form's own label
    rather than on top of it."""
    import docx
    from docx.oxml.ns import qn

    for n in ("Main Student", "Main Helper", "Aux Student", "Aux Helper",
              "Counselor Man"):
        core.add_student(n, "Brother", core.PRIVILEGES)
    ids = dict(zip(core.get_students()["name"], core.get_students()["id"]))
    names = {v: k for k, v in ids.items()}
    slots = core.apply_aux(core.build_midweek_slots(core.default_midweek_parts()),
                           True)
    picks = {}
    for i, s in enumerate(slots):
        if s["role"] == "Aux Classroom Counselor":
            picks[i] = (ids["Counselor Man"], None)
        elif s["student_part"]:
            aux = s["hall"] != core.MAIN_HALL
            picks[i] = (ids["Aux Student" if aux else "Main Student"],
                        ids["Aux Helper" if aux else "Main Helper"]
                        if s["needs_assistant"] else None)
        else:
            picks[i] = (people["Kofi Mensah"], None)
    core.save_schedule("2026-09-16", core.MIDWEEK, slots, picks,
                       {"heading": "SEPTEMBER 14-20", "aux": True, "aux_group": "1",
                        "opening_song": "Song 74"}, names)
    data, skipped = core.build_s140_data([("2026-09-16", core.MIDWEEK)],
                                         core.get_schedules(), "Teshie Asafo",
                                         "GROUP")
    assert not skipped and data["aux"]

    def txt(el):
        return "".join(n.text or "" for n in el.iter(qn("w:t")))

    out = core.fill_s140(S140_GA.read_bytes(), data)
    trs = docx.Document(io.BytesIO(out)).tables[0]._tbl.findall(qn("w:tr"))
    rows = [[txt(c) for c in tr.findall(qn("w:tc"))] for tr in trs]

    # the counselor sits beside the label, which survives
    assert any(r[1:3] == ["Asa 2 Ŋaawolɔ:", "Counselor Man"] for r in rows if len(r) > 2)
    # classroom names in the Asa 2 column, main hall in Asa 1
    paired = [r for r in rows if len(r) >= 5 and "Aux Student" in r[3]]
    assert paired, rows
    for r in paired:
        assert r[3].startswith("Aux Student") and r[4].startswith("Main Student")
    # and the column captions are kept for a classroom week
    assert any("Asa 2" in r[1] and "Asa 1" in r[2] for r in rows if len(r) > 2)


@pytest.mark.skipif(not S140_GA.is_file(), reason="S-140 blanks not in the repo")
def test_s140_top_line_and_group(core, people):
    """The blank leaves two empty rows above the date. They carry the meeting
    and the congregation, and the classroom's group goes on the counselor's
    line — all in the form's own font, not Word's default."""
    import docx
    from docx.oxml.ns import qn

    for n in ("Aux Student", "Counselor Man"):
        core.add_student(n, "Brother", core.PRIVILEGES)
    ids = dict(zip(core.get_students()["name"], core.get_students()["id"]))
    names = {v: k for k, v in ids.items()}
    slots = core.apply_aux(core.build_midweek_slots(core.default_midweek_parts()),
                           True)
    picks = {}
    for i, s in enumerate(slots):
        if s["role"] == "Aux Classroom Counselor":
            picks[i] = (ids["Counselor Man"], None)
        elif s["student_part"]:
            picks[i] = (ids["Aux Student"], None)
        else:
            picks[i] = (people["Kofi Mensah"], None)
    core.save_schedule("2026-09-16", core.MIDWEEK, slots, picks,
                       {"heading": "SEPTEMBER 14-20", "aux": True,
                        "aux_group": "1", "opening_song": "Song 74"}, names)
    data, _ = core.build_s140_data([("2026-09-16", core.MIDWEEK)],
                                   core.get_schedules(), "Teshie Asafo", "GROUP")
    assert data["weeks"][0]["aux_group"] == "1"
    data["meeting_name"] = core.TRANSLATIONS["Ga"]["midweek_meeting"]

    def txt(el):
        return "".join(n.text or "" for n in el.iter(qn("w:t")))

    trs = docx.Document(io.BytesIO(
        core.fill_s140(S140_GA.read_bytes(), data))).tables[0]._tbl.findall(qn("w:tr"))
    rows = [[txt(c) for c in tr.findall(qn("w:tc"))] for tr in trs]

    assert rows[0][0] == "Wɔshiɛmɔ Kɛ Wɔshihilɛ Kpee"
    assert rows[0][1] == "Teshie Asafo"
    # heading the document once, not once per week
    assert sum(1 for r in rows if "Wɔshiɛmɔ Kɛ Wɔshihilɛ" in " ".join(r)) == 1
    # at the date's size and weight, not the small grey label's
    title_run = next(r for r in trs[0].findall(qn("w:tc"))[0].iter(qn("w:r")))
    props = {c.tag.split("}")[1] for c in title_run.find(qn("w:rPr"))}
    assert "b" in props and "sz" not in props
    # the group shares the counselor's line and does not repeat the room name
    counselor = next(r for r in rows if len(r) > 2 and r[1] == "Asa 2 Ŋaawolɔ:")
    assert counselor[0] == "Kuu 1" and counselor[2] == "Counselor Man"

    # and the added text carries run formatting rather than Word's default
    top_cells = trs[0].findall(qn("w:tc"))
    for cell in top_cells[:2]:
        run = next(r for r in cell.iter(qn("w:r")))
        assert run.find(qn("w:rPr")) is not None


def test_unfilled_parts_print_a_dash(core, people):
    """An empty cell on a noticeboard reads as a fault in the sheet; a dash
    reads as a part still to be filled."""
    slots = core.build_midweek_slots(core.default_midweek_parts())
    names = dict(zip(core.get_students()["id"], core.get_students()["name"]))
    chairman = next(i for i, s in enumerate(slots) if s["role"] == "Chairman")
    core.save_schedule("2026-09-16", core.MIDWEEK, slots,
                       {chairman: (people["Kofi Mensah"], None)}, {}, names)
    text = _text(core.generate_schedule_pdf([("2026-09-16", core.MIDWEEK)],
                                            core.get_schedules()))
    assert "Kofi Mensah" in text
    assert "\u2014" in text                      # the unfilled parts


def test_weeks_flow_onto_a_sheet_rather_than_one_each(core, people):
    """Four weekend weeks share one A4. Grouping every meeting separately for
    the two-up work put a page break between each of them instead."""
    import io

    import pypdf

    names = dict(zip(core.get_students()["id"], core.get_students()["name"]))
    dates = ["2026-10-03", "2026-10-10", "2026-10-17", "2026-10-24"]
    for d in dates:
        slots = core.default_weekend_slots()
        core.save_schedule(d, core.WEEKEND, slots,
                           {i: (people["Kofi Mensah"], None)
                            for i in range(len(slots))},
                           {"talk_number": "73", "talk_title": "A title"}, names)
    rows = core.get_schedules()
    meetings = [(d, core.WEEKEND) for d in dates]

    pdf = core.generate_schedule_pdf(meetings, rows)
    assert len(pypdf.PdfReader(io.BytesIO(pdf)).pages) == 1
    # and the meeting name heads the sheet once, not once per week
    assert _text(pdf).count("Weekend Meeting") == 1
